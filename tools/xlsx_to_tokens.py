#!/usr/bin/env python3
"""Generate _shared/tokens.css from design/MASTER-STYLES.xlsx.

The Text, Colour, Spacing and Grid sheets are emitted. Box, Components and
Top bar are read and emit nothing. ChangeLog is ignored. Grid rows marked
LEGACY-PR1 are not emitted.

Day (the :root block) is the default. Dark is the override, in the same
two-block shape as main: prefers-color-scheme, then [data-theme="dark"].
Desktop type-scale values live on :root of tokens.css and tokens-core.css.
The four mobile type-scale declarations stay in the tokens.css 560 block and
are mirrored in type-scale.css. tokens-core.css is that file minus exactly
those four lines. --check fails if the copies differ from the workbook.

Usage:
  python3 tools/xlsx_to_tokens.py            # write tokens.css and tokens-core.css
  python3 tools/xlsx_to_tokens.py --check    # MATCH or a diff, exit 1 on drift
  python3 tools/xlsx_to_tokens.py --owned    # report-only baseline comparison
  python3 tools/xlsx_to_tokens.py --map-lint # report-only hardcoded-value counts
  python3 tools/xlsx_to_tokens.py --map-lint-shared
  python3 tools/xlsx_to_tokens.py --settings-cover
  python3 tools/xlsx_to_tokens.py --repo-grep
"""
import difflib
import html
import re
import subprocess
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "design" / "MASTER-STYLES.xlsx"
TOKENS = ROOT / "_shared" / "tokens.css"
TOKENS_CORE = ROOT / "_shared" / "tokens-core.css"
TYPE_SCALE = ROOT / "_shared" / "type-scale.css"
# PR3 re-baselines "existing token bytes kept" on this commit, not origin/main.
TOKEN_BASE = "03ca4a576eff907e72ec113cb8ae27142f09ff84"
CARD_RADIUS_BASE = "8px"
CARD_RADIUS_HEAD = ".5rem"
DESKTOP_SCALE = {
    "--type-scale": "1",
    "--vis-body": "12px",
    "--vis-meta": "11px",
    "--nx-chip-font": "10px",
}
SCALE_NAMES = tuple(DESKTOP_SCALE)
# Shown on /settings already. New PR3 tokens stay out of that page so it
# stays byte-identical (stamp-only) and pixel-identical.
SETTINGS_HIDE_PREFIXES = ("LEGACY-PR3", "NEW PR3")
# These names are new in PR3. /settings already lists every older token.
SETTINGS_HIDE_NAMES = {
    "--font-mono",
    "--cite-font",
    "--radius-m",
    "--border-see",
    "--border-na",
}
NEW_ROOT_NAMES = (
    "--font-mono",
    "--cite-font",
    "--radius-m",
    "--border-see",
    "--border-na",
)
DESKTOP_DECLS = (
    "--type-scale: 1;",
    "--vis-body: 12px;",
    "--vis-meta: 11px;",
    "--nx-chip-font: 10px;",
)
# Indent 4, first 560 block only. .is-mobile copies use indent 2.
CORE_DROP_LINES = (
    "    --type-scale: 0.85;",
    "    --vis-body: calc(12px * var(--type-scale));",
    "    --vis-meta: calc(11px * var(--type-scale));",
    "    --nx-chip-font: calc(10px * var(--type-scale));",
)
GUARD_MARK = (
    "/* Settings keeps the pre-PR3 desktop fallback for these four tokens."
)
SETTINGS_GUARD = """/* Settings keeps the pre-PR3 desktop fallback for these four tokens.
   :root defines them for every other page. Mobile still follows the 560 block. */
body[data-pack="site"] {
  --type-scale: initial;
  --vis-body: initial;
  --vis-meta: initial;
  --nx-chip-font: initial;
}
@media (max-width:560px) {
  body[data-pack="site"] {
    --type-scale: 0.85;
    --vis-body: calc(12px * var(--type-scale));
    --vis-meta: calc(11px * var(--type-scale));
    --nx-chip-font: calc(10px * var(--type-scale));
  }
}
"""
SETTINGS_CSS = ROOT / "_shared" / "settings.css"
SETTINGS_ROWS = ROOT / "_shared" / "settings-rows.html"
SETTINGS_GEN = ROOT / "_shared" / "settings-gen.css"
SETTINGS_PAGE = ROOT / "site" / "settings" / "index.html"
ROW_BEGIN = "<!-- BEGIN SETTINGS ROWS -->"
ROW_END = "<!-- END SETTINGS ROWS -->"
# Appended after the existing :root declarations. Order is the brief's.
NEW_GRID_ORDER = (
    "--preview-w-mobile",
    "--preview-w-tablet",
    "--preview-w-desktop",
    "--bp-mobile-n",
    "--bp-wide-n",
    "--preview-w-mobile-n",
    "--preview-w-tablet-n",
    "--preview-w-desktop-n",
    "--ruler-max-n",
)
NEW_GRID = set(NEW_GRID_ORDER)
# Each unitless copy must equal this length token, in px (16px per rem).
UNITLESS_PX = (
    ("--bp-mobile-n", "--bp-mobile"),
    ("--bp-wide-n", "--bp-wide"),
    ("--preview-w-mobile-n", "--preview-w-mobile"),
    ("--preview-w-tablet-n", "--preview-w-tablet"),
    ("--preview-w-desktop-n", "--preview-w-desktop"),
    ("--ruler-max-n", "--preview-w-desktop"),
)
SAMPLE = "Inventories are measured at the lower of cost and NRV"
INK_SKIP = {"--bg", "--surface", "--line", "--band", "--nest"}
RULER_MARKS = (
    ("--preview-w-mobile-n", "390", False),
    ("--bp-mobile-n", "560", True),
    ("--bp-wide-n", "800", True),
    ("--preview-w-tablet-n", "922", False),
    ("--preview-w-desktop-n", "1280", False),
)
READ_SHEETS = ("Text", "Colour", "Box", "Spacing", "Grid", "Components", "Top bar")
EXTRA = re.compile(r"(\w+)=([\d.]+rem)/([\d.]+rem)")
DECL = re.compile(r"(--[\w-]+)\s*:\s*([^;]+);")
PROP = re.compile(r"(--[\w-]+)\s*:")
TOKEN_NAME = re.compile(r"--[\w-]+")


def n(value):
    """rem number format from the box draft: no leading zero."""
    text = "%g" % float(value)
    return text[1:] if text.startswith("0.") else text


def rem(value):
    return n(value) + "rem"


def line(name, value, indent):
    return (" " * indent) + f"{name}: {value};"


def extras(notes):
    return EXTRA.findall(str(notes or ""))


def css_scalar(value):
    """Cell value to CSS text. Whole numbers stay unitless integers."""
    if isinstance(value, bool):
        raise SystemExit(f"unexpected boolean token value {value!r}")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def length_px(value):
    text = str(value).strip()
    if text.endswith("px"):
        return float(text[:-2])
    if text.endswith("rem"):
        return float(text[:-3]) * 16
    return float(text)


def load_workbook():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    missing = [name for name in READ_SHEETS if name not in wb.sheetnames]
    if missing:
        raise SystemExit("missing sheet: " + ", ".join(missing))
    # Read, then emit nothing. ChangeLog is ignored.
    for name in ("Box", "Components", "Top bar"):
        wb[name].max_row
    return wb


def sheet_rows(wb, name):
    ws = wb[name]
    header = [cell.value for cell in ws[1]]
    rows = []
    for cells in ws.iter_rows(min_row=2):
        if cells[0].value is None and all(cell.value is None for cell in cells):
            continue
        rows.append(dict(zip(header, [cell.value for cell in cells])))
    return rows


def _type_scale_rows(text_rows):
    rows = [row for row in text_rows if row.get("Group") == "type-scale"]
    if len(rows) != 4:
        raise SystemExit(f"expected 4 type-scale rows, found {len(rows)}")
    return rows


def type_scale_mobile_value(row):
    mobile_px = row.get("Mobile px")
    if mobile_px is not None and str(mobile_px).strip() != "":
        return str(mobile_px).strip()
    notes = str(row.get("Notes") or "")
    if "mobile " not in notes:
        raise SystemExit(f"type-scale row {row.get('Token prefix')} has no mobile value")
    return notes.split("mobile ", 1)[1].strip()


def type_scale_lines(text_rows):
    """Mobile mirror only. Desktop values are a separate :root block."""
    lines = []
    for row in _type_scale_rows(text_rows):
        lines.append(line(row["Token prefix"], type_scale_mobile_value(row), 4))
    return lines


def type_scale_desktop_lines(text_rows):
    lines = []
    for row in _type_scale_rows(text_rows):
        desktop = row.get("Desktop px")
        if desktop is None or str(desktop).strip() == "":
            raise SystemExit(f"type-scale row {row.get('Token prefix')} has no desktop value")
        lines.append(line(row["Token prefix"], str(desktop).strip(), 2))
    return lines


def emit_map_row(row, bucket, indent, mobile):
    prefix = row["Token prefix"]
    if mobile:
        if float(row["Mobile rem"]) != float(row["Desktop rem"]):
            bucket.append(line(f"{prefix}-fs", rem(row["Mobile rem"]), indent))
        for key, desktop, mobile_value in extras(row.get("Notes")):
            if desktop != mobile_value:
                bucket.append(line(f"{prefix}-{key}", mobile_value, indent))
        return
    bucket.append(line(f"{prefix}-fs", rem(row["Desktop rem"]), indent))
    if row["Weight"] is not None:
        bucket.append(line(f"{prefix}-fw", str(int(row["Weight"])), indent))
    if row["Line height"] is not None:
        bucket.append(line(f"{prefix}-lh", n(row["Line height"]), indent))
    for key, desktop, _mobile_value in extras(row.get("Notes")):
        bucket.append(line(f"{prefix}-{key}", desktop, indent))


def emit_teaching_row(row, bucket):
    prefix = row["Token prefix"]
    if prefix == "--cite" or row.get("Style name") == "Cite":
        bucket.append(line("--cite-size", rem(row["Desktop rem"]), 2))
        bucket.append(line("--cite-weight", str(int(row["Weight"])), 2))
        return
    emit_map_row(row, bucket, 2, mobile=False)


def build_parts(wb):
    text_rows = sheet_rows(wb, "Text")
    colour_rows = sheet_rows(wb, "Colour")
    spacing_rows = sheet_rows(wb, "Spacing")
    grid_rows = sheet_rows(wb, "Grid")

    day = []
    colour_day = []
    map_mobile = []
    grid_mobile = []
    night = []

    for row in colour_rows:
        token = row.get("Token")
        if not token or not str(token).startswith("--"):
            continue
        day_line = line(token, str(row["Day value (default :root)"]).strip(), 2)
        day.append(day_line)
        colour_day.append(day_line)
        night_value = row.get("Night value (dark override: prefers-color-scheme + [data-theme=dark])")
        if night_value is None or str(night_value).strip() == "(same as day)":
            continue
        night.append(line(token, str(night_value).strip(), 4))

    for row in text_rows:
        if row.get("Group") != "site":
            continue
        day.append(line(row["Token prefix"], rem(row["Desktop rem"]), 2))

    for row in spacing_rows:
        token = row.get("Token")
        if not token or not str(token).startswith("--"):
            continue
        if row.get("Desktop") is None:
            continue
        day.append(line(str(token), str(row["Desktop"]).strip(), 2))

    for row in grid_rows:
        if row.get("Status") == "LEGACY-PR1":
            continue
        if row.get("Status") == "GENERATED-MQ":
            continue
        token = row.get("Token")
        if not token or not str(token).startswith("--"):
            continue
        if str(token) in NEW_GRID:
            continue
        if row.get("Desktop") is None:
            continue
        day.append(line(str(token), css_scalar(row["Desktop"]), 2))
        mobile_value = row.get("Mobile")
        if mobile_value is None:
            continue
        mobile_text = str(mobile_value).strip()
        if mobile_text in ("", "same"):
            continue
        if mobile_text != str(row["Desktop"]).strip():
            grid_mobile.append(line(str(token), mobile_text, 4))

    for row in text_rows:
        if row.get("Group") != "map":
            continue
        emit_map_row(row, day, 2, mobile=False)
        emit_map_row(row, map_mobile, 4, mobile=True)

    for row in text_rows:
        if row.get("Group") != "teaching":
            continue
        emit_teaching_row(row, day)

    scale = type_scale_lines(text_rows)
    # Text sheet (map overrides, then the type-scale mirror), then Grid.
    mobile = map_mobile + scale + grid_mobile

    by_token = {}
    for row in grid_rows:
        token = row.get("Token")
        if token:
            by_token[str(token)] = row
    for name in NEW_GRID_ORDER:
        row = by_token.get(name)
        if row is None or row.get("Desktop") is None:
            raise SystemExit(f"missing grid token {name}")
        day.append(line(name, css_scalar(row["Desktop"]), 2))
    font_rows = [row for row in text_rows if row.get("Group") == "site-font"]
    if not font_rows:
        raise SystemExit("expected at least 1 site-font row")
    for row in font_rows:
        font_value = str(row.get("Notes") or "").strip()
        if not font_value:
            raise SystemExit(f"site-font row {row.get('Token prefix')} has no value in Notes")
        day.append(line(str(row["Token prefix"]), font_value, 2))
    day.extend(type_scale_desktop_lines(text_rows))

    if len(night) != 43:
        raise SystemExit(f"expected 43 night colours, found {len(night)}")
    if len(colour_day) != 43:
        raise SystemExit(f"expected 43 day colours, found {len(colour_day)}")
    return day, mobile, night, scale, colour_day


def block(selector, decls, indent):
    pad = " " * indent
    body = "\n".join(decls)
    return f"{pad}{selector} {{\n{body}\n{pad}}}"


def render(day, mobile, night, colour_day):
    # Night declarations sit at indent 4 inside the media :root and at indent 2
    # in the data-theme block. build_parts stored them at indent 4; restamp
    # the data-theme copy at indent 2 so both blocks carry the same values.
    night_inner = night
    night_root = [line(name, value, 2) for name, value in parse_decl_lines(night)]
    mobile_root = [line(name, value, 2) for name, value in parse_decl_lines(mobile)]
    parts = [
        "/* GENERATED — do not edit. tools/xlsx_to_tokens.py reads design/MASTER-STYLES.xlsx. */",
        block(":root", day, 0),
        "@media (max-width:560px) {\n" + block(":root", mobile, 2) + "\n}",
        "@media (prefers-color-scheme: dark) {\n"
        + block(':root:not([data-theme="light"])', night_inner, 2)
        + "\n}",
        block(':root[data-theme="dark"]', night_root, 0),
        block(".is-mobile", mobile_root, 0),
        "/* settings preview only - scoped dark */\n"
        + block('.settings-preview[data-theme="dark"]', night_root, 0),
        "/* settings preview only - scoped light */\n"
        + block('.settings-preview[data-theme="light"]', colour_day, 0),
    ]
    return "\n".join(parts) + "\n" + SETTINGS_GUARD


def render_core(css):
    """tokens.css minus the four mobile type-scale lines in the first 560 block."""
    start = css.index("@media (max-width:560px)")
    end = css.index("@media (prefers-color-scheme", start)
    block_text = css[start:end]
    for item in CORE_DROP_LINES:
        needle = item + "\n"
        count = block_text.count(needle)
        if count != 1:
            raise SystemExit(
                f"tokens-core: expected one {item!r} in the first 560 block, found {count}"
            )
        block_text = block_text.replace(needle, "", 1)
    return css[:start] + block_text + css[end:]


def core_removals(full, core):
    """Line-level deletions of core versus full. Anything but a pure delete is an error."""
    full_lines = full.splitlines()
    core_lines = core.splitlines()
    removed = []
    matcher = difflib.SequenceMatcher(a=full_lines, b=core_lines)
    for tag, i1, i2, _j1, _j2 in matcher.get_opcodes():
        if tag == "equal":
            continue
        if tag != "delete":
            return None
        removed.extend(full_lines[i1:i2])
    return removed


def git_at(rev, rel):
    try:
        return subprocess.check_output(
            ["git", "show", f"{rev}:{rel}"],
            cwd=ROOT,
            text=True,
            stderr=subprocess.DEVNULL,
        )
    except subprocess.CalledProcessError:
        raise SystemExit(f"cannot read {rel} from {rev}")


def parse_decl_lines(lines):
    found = []
    for item in lines:
        match = DECL.search(item)
        if not match:
            raise SystemExit(f"declaration did not parse: {item}")
        found.append((match.group(1), match.group(2).strip()))
    return found


def git_text(rel):
    import os
    refs = ["origin/main", "main"]
    base = os.environ.get("GITHUB_BASE_REF", "")
    if base:
        refs.extend([f"origin/{base}", base])
    for ref in refs:
        try:
            return subprocess.check_output(
                ["git", "show", f"{ref}:{rel}"],
                cwd=ROOT,
                text=True,
                stderr=subprocess.DEVNULL,
            )
        except subprocess.CalledProcessError:
            continue
    raise SystemExit(f"cannot read {rel} from main")


def slice_blocks(css):
    """Split tokens.css. The 560 block is not part of :root.

    The old slice ran from `:root {` to the prefers-color-scheme block, so
    the 560 declarations were counted as day values and baseline_stop
    refused a correct file. Day ends at the first @media.
    """
    root_at = css.index(":root {")
    narrow_at = css.index("@media (max-width:560px)")
    media_at = css.index("@media (prefers-color-scheme")
    dark_at = css.index(':root[data-theme="dark"]')
    dark_end = len(css)
    for marker in ("\n.is-mobile {", "\n/* settings preview"):
        at = css.find(marker, dark_at)
        if at != -1:
            dark_end = min(dark_end, at)
    return {
        "root": DECL.findall(css[root_at:narrow_at]),
        "narrow": DECL.findall(css[narrow_at:media_at]),
        "mq": DECL.findall(css[media_at:dark_at]),
        "dark": DECL.findall(css[dark_at:dark_end]),
    }


def baseline_stop(day, mobile, night, scale):
    """STOP if a base-commit token would change value. Additions are allowed."""
    base = slice_blocks(git_at(TOKEN_BASE, "_shared/tokens.css"))
    generated_root = dict(parse_decl_lines(day))
    generated_mobile = dict(parse_decl_lines(mobile))
    generated_night = dict(parse_decl_lines(night))
    problems = []
    for name, value in base["root"]:
        got = generated_root.get(name)
        if (
            name == "--card-radius"
            and value.strip() == CARD_RADIUS_BASE
            and (got or "").strip() == CARD_RADIUS_HEAD
        ):
            continue
        if got is None or got.strip() != value.strip():
            problems.append(f":root {name} base={value!r} generated={got!r}")
    for name, value in base["narrow"]:
        got = generated_mobile.get(name)
        if got is None or got.strip() != value.strip():
            problems.append(f"560 {name} base={value!r} generated={got!r}")
    for block_name in ("mq", "dark"):
        for name, value in base[block_name]:
            got = generated_night.get(name)
            if got is None or got.strip() != value.strip():
                problems.append(f"{block_name} {name} base={value!r} generated={got!r}")
    current = type_scale_file_text()
    expected = "\n".join(scale)
    if expected not in current:
        problems.append("type-scale.css 4-token block is not the verbatim mirror")
    if problems:
        print("STOP: token value would change, or the type-scale mirror does not match.")
        for item in problems:
            print(item)
        raise SystemExit(1)


def type_scale_file_text():
    """Entry plus the module that holds the mobile block after the split."""
    parts = []
    module = ROOT / "_shared" / "modules" / "type-scale" / "text-1.css"
    for path in (TYPE_SCALE, module):
        if path.exists():
            parts.append(path.read_text(encoding="utf-8"))
    return "\n".join(parts)


def type_scale_mirror_ok(css, scale):
    expected = "\n".join(scale)
    file_text = type_scale_file_text()
    in_tokens = expected in css
    in_file = expected in file_text
    return in_tokens and in_file, expected, file_text


def existing_tokens_kept(css):
    """Bytes at TOKEN_BASE, plus only the PR3 rows in section 3.

    The comparison is `git show <base>:_shared/tokens.css`, not origin/main.
    Allowed edits are --card-radius 8px to .5rem, the new :root rows, and the
    settings-page fallback appended after the generated blocks.
    """
    base = git_at(TOKEN_BASE, "_shared/tokens.css")
    head = css
    mark = head.find(GUARD_MARK)
    if mark < 0:
        return False, "missing settings fallback block"
    head = head[:mark]
    if not head.endswith("\n"):
        head += "\n"
    kept = []
    for line in head.splitlines(keepends=True):
        stripped = line.strip()
        if stripped.startswith("--legacy-") and stripped.endswith(";"):
            continue
        if any(stripped.startswith(name + ":") for name in NEW_ROOT_NAMES):
            continue
        if stripped in DESKTOP_DECLS:
            continue
        if stripped == "--card-radius: .5rem;":
            line = line.replace("--card-radius: .5rem;", "--card-radius: 8px;", 1)
        kept.append(line)
    folded = "".join(kept)
    if folded != base:
        diff = "".join(difflib.unified_diff(
            base.splitlines(keepends=True),
            folded.splitlines(keepends=True),
            fromfile=f"{TOKEN_BASE}:_shared/tokens.css",
            tofile="kept",
        ))
        return False, diff
    return True, ""


def render_settings_gen(grid_rows):
    """One 560 block. C3: no 800 block, even if a workbook row still lists one."""
    rows = [row for row in grid_rows if row.get("Status") == "GENERATED-MQ"]
    if len(rows) != 1:
        raise SystemExit(f"expected 1 GENERATED-MQ row, found {len(rows)}")
    desktop = "" if rows[0].get("Desktop") is None else css_scalar(rows[0]["Desktop"])
    mobile = rows[0].get("Mobile")
    mobile_text = "" if mobile is None else css_scalar(mobile)
    if "800" in desktop or "800" in mobile_text:
        print(
            "R1: GENERATED-MQ row still lists an 800 block. "
            "C3 emits @media (max-width:560px) only."
        )
    elif desktop != "@media (max-width:560px)":
        raise SystemExit(f"GENERATED-MQ desktop is {desktop!r}, expected the 560 block")
    return """/* GENERATED - do not edit. tools/xlsx_to_tokens.py reads design/MASTER-STYLES.xlsx. */
@media (max-width:560px) {
  .toggles .seg button[data-set="m"] {
    background: var(--accent-soft);
    color: var(--accent);
  }
  .toggles .seg button[data-set="d"] {
    background: var(--surface);
    color: var(--muted);
  }
  .toggles .state-d { display: none; }
  .toggles .state-m { display: inline; }
  .settings-preview .now {
    inline-size: calc(var(--preview-w-mobile-n) / var(--ruler-max-n) * 100%);
  }
  .settings-preview .pv {
    max-inline-size: var(--preview-w-mobile);
    outline: var(--stroke-1) solid var(--line);
    outline-offset: var(--space-1);
  }
  .settings-preview .pv:has(.ruler) {
    max-inline-size: none;
    outline-style: none;
  }
  .settings-preview .v-m,
  body.mode-d .settings-preview .v-m {
    color: var(--accent);
    border-color: var(--accent);
    background: var(--accent-soft);
  }
  .settings-preview .v-d,
  body.mode-d .settings-preview .v-d {
    color: var(--muted);
    background: var(--surface);
    border-color: var(--surface);
  }
  .settings-preview .set-table,
  .settings-preview .set-table tbody,
  .settings-preview .set-table tr,
  .settings-preview .set-table td,
  .settings-preview .set-table th {
    display: block;
    inline-size: auto;
  }
  .settings-preview .set-table thead { display: none; }
  .settings-preview .set-table tr { margin-bottom: var(--space-cell); }
  .settings-preview .row td.l {
    border-end-start-radius: 0;
    border-end-end-radius: 0;
  }
  .settings-preview .row td.r {
    border-block-start: 0;
    border-start-start-radius: 0;
    border-start-end-radius: 0;
  }
}
"""


def _pair(name, root_map, mobile_map):
    desktop = root_map[name]
    mobile = mobile_map.get(name)
    if mobile is None or mobile == desktop:
        return name, desktop, "="
    return name, desktop, mobile


def _token_text(text):
    """Escape displayed text. A token name stays one word, hyphen included."""
    return TOKEN_NAME.sub(
        lambda match: f'<span class="token-name">{match.group(0)}</span>',
        html.escape(str(text)),
    )


def _code_name(name):
    return f'<span class="token-name"><code>{html.escape(name)}</code></span>'


def _vals(name, desktop, mobile):
    return (
        f'<div class="tok">{_code_name(name)}'
        f'<span class="v v-d"><b>D</b> {_token_text(desktop)}</span>'
        f'<span class="v v-m"><b>M</b> {_token_text(mobile)}</span></div>'
    )


def _section(title, kind=""):
    extra = f" {kind}" if kind else ""
    return f'<tr class="sec{extra}"><th colspan="2" scope="rowgroup">{html.escape(title)}</th></tr>'


def _data_row(kind, left, preview):
    return (
        f'<tr class="row {kind}"><td class="l">{left}</td>'
        f'<td class="r"><div class="pv">{preview}</div></td></tr>'
    )


def _ruler():
    marks = []
    for token, label, solid in RULER_MARKS:
        klass = "mk tok" if solid else "mk"
        marks.append(
            f'<span class="{klass}" style="--p-x:var({token})"><i></i><em>{html.escape(label)}</em></span>'
        )
    return f'<div class="ruler">{"".join(marks)}<span class="now"></span></div>'


def _grid_preview(token):
    if token.endswith("-n") or token in ("--bp-mobile", "--bp-wide"):
        unit = {"--bp-mobile": "--bp-mobile-n", "--bp-wide": "--bp-wide-n"}.get(token, token)
        solid = " tok" if unit in ("--bp-mobile-n", "--bp-wide-n") else ""
        return (
            f'<div class="ruler"><span class="mk{solid}" style="--p-x:var({unit})"><i></i></span>'
            f'<span class="now"></span></div>'
        )
    if token.startswith("--cols-"):
        count = int(token[-1])
        return f'<div class="cols" style="--p-cols:var({token})">{"<span></span>" * count}</div>'
    if token.startswith("--rail-bar"):
        return (
            f'<div class="rail" style="--p-w:var({token})"><span class="rail-bar"></span>'
            f'<span class="rail-body">Rail detail</span></div>'
        )
    if token.startswith("--stroke"):
        return f'<span class="stroke" style="--p-w:var({token})"></span>'
    if token.startswith("--preview-w"):
        return f'<span class="bar pw" style="--p-w:var({token})"></span>'
    if token == "--page-max":
        return _ruler() + '<span class="page-max-bar"></span>'
    return f'<span class="bar" style="--p-w:var({token})"></span>'


def settings_rows_text(wb, day, mobile):
    """One row per style. Every token name is in exactly one row's code element."""
    root_map = dict(parse_decl_lines(day))
    mobile_map = dict(parse_decl_lines(mobile))
    text_rows = sheet_rows(wb, "Text")
    colour_rows = sheet_rows(wb, "Colour")
    spacing_rows = sheet_rows(wb, "Spacing")
    grid_rows = sheet_rows(wb, "Grid")
    out = [_section("Text")]
    for row in text_rows:
        prefix = row.get("Token prefix")
        if not prefix:
            continue
        group = row.get("Group") or ""
        style = row.get("Style name") or prefix
        role = row.get("Role / where used") or ""
        if group == "site-font" and str(prefix) != "--font-family":
            continue
        if group == "type-scale":
            tokens = [(prefix, "fallback 1 / unset", mobile_map[prefix])]
            style_attr = (
                "--p-fs:calc(var(--t-body) * var(--type-scale))"
                if prefix == "--type-scale"
                else f"--p-fs:var({prefix})"
            )
        elif group == "site-font":
            tokens = [_pair(prefix, root_map, mobile_map)]
            style_attr = "--p-family:var(--font-family)"
        elif group == "site":
            tokens = [_pair(prefix, root_map, mobile_map)]
            style_attr = f"--p-fs:var({prefix})"
        elif prefix == "--cite" or style == "Cite":
            tokens = [
                _pair("--cite-size", root_map, mobile_map),
                _pair("--cite-weight", root_map, mobile_map),
            ]
            style_attr = "--p-fs:var(--cite-size);--p-fw:var(--cite-weight)"
        else:
            tokens = [_pair(f"{prefix}-fs", root_map, mobile_map)]
            parts = [f"--p-fs:var({prefix}-fs)"]
            if row.get("Weight") is not None:
                tokens.append(_pair(f"{prefix}-fw", root_map, mobile_map))
                parts.append(f"--p-fw:var({prefix}-fw)")
            if row.get("Line height") is not None:
                tokens.append(_pair(f"{prefix}-lh", root_map, mobile_map))
                parts.append(f"--p-lh:var({prefix}-lh)")
            if prefix == "--mchip":
                for key, _desktop, _mobile in extras(row.get("Notes")):
                    tokens.append(_pair(f"{prefix}-{key}", root_map, mobile_map))
            style_attr = ";".join(parts)
        left = (
            f'<div class="meta"><div class="sty">{html.escape(str(style))} '
            f'<span class="grp">{html.escape(str(group))}</span></div>'
            + "".join(_vals(*item) for item in tokens)
            + f'<div class="role">{_token_text(role)}</div></div>'
        )
        preview = f'<p class="sample" style="{style_attr}">{SAMPLE}</p>'
        out.append(_data_row("text", left, preview))
    out.append(_section("Colour"))
    for row in colour_rows:
        token = row.get("Token")
        if not token or not str(token).startswith("--"):
            continue
        token = str(token)
        day_value = root_map[token]
        night_value = row.get("Night value (dark override: prefers-color-scheme + [data-theme=dark])")
        night_text = day_value if night_value is None or str(night_value).strip() == "(same as day)" else str(night_value).strip()
        role = row.get("Role") or ""
        left = (
            f'<div class="meta"><div class="tok">{_code_name(token)}'
            f'<span class="v v-day"><b>Day</b> {_token_text(day_value)}</span>'
            f'<span class="v v-night"><b>Night</b> {_token_text(night_text)}</span></div>'
            f'<div class="role">{_token_text(role)}</div></div>'
        )
        chip = ""
        if not token.endswith("-soft") and token not in INK_SKIP:
            chip = f'<span class="chip" style="--p-c:var({token})">Aa text</span>'
        preview = f'<span class="sw" style="--p-c:var({token})"></span>{chip}'
        out.append(_data_row("colour", left, preview))
    out.append(_section("Spacing"))
    for row in spacing_rows:
        token = row.get("Token")
        if not token or not str(token).startswith("--"):
            continue
        token = str(token)
        role = row.get("Use") or ""
        if role.startswith(SETTINGS_HIDE_PREFIXES) or token.startswith("--legacy-"):
            continue
        if token in SETTINGS_HIDE_NAMES:
            continue
        name, desktop, mobile_value = _pair(token, root_map, mobile_map)
        if token == "--card-radius":
            # The stylesheet value is .5rem. The settings cell stays 8px so the
            # page HTML does not change.
            desktop = CARD_RADIUS_BASE
            mobile_value = "="
            preview = f'<span class="radius" style="--p-r:var({token})"></span>'
        else:
            preview = f'<span class="bar" style="--p-w:var({token})"></span>'
        left = (
            f'<div class="meta">{_vals(name, desktop, mobile_value)}'
            f'<div class="role">{_token_text(role)}</div></div>'
        )
        out.append(_data_row("spacing", left, preview))
    out.append(_section("Grid"))
    for row in grid_rows:
        if row.get("Status") in ("LEGACY-PR1", "GENERATED-MQ"):
            continue
        token = row.get("Token")
        if not token or not str(token).startswith("--"):
            continue
        token = str(token)
        name, desktop, mobile_value = _pair(token, root_map, mobile_map)
        role = row.get("Use / current usage") or ""
        left = (
            f'<div class="meta">{_vals(name, desktop, mobile_value)}'
            f'<div class="role">{_token_text(role)}</div></div>'
        )
        out.append(_data_row("grid", left, _grid_preview(token)))
    ruler_left = (
        '<div class="meta"><div class="tok"><code>breakpoint ruler</code></div>'
        '<div class="role">'
        + _token_text(
            "390 / 560 / 800 / 922 / 1280 - solid marks are "
            "--bp-mobile and --bp-wide, hollow marks are preview widths"
        )
        + "</div></div>"
    )
    out.append(_data_row("grid ruler-row", ruler_left, _ruler()))
    for title in ("Box", "Components", "TopBar"):
        out.append(_section(title, "next"))
        out.append(
            '<tr class="row next"><td class="l" colspan="2"><div class="role">'
            'Next phase - tokens not yet defined in MASTER-STYLES.xlsx; no preview.'
            '</div></td></tr>'
        )
    return "\n".join(out) + "\n"


def inject_rows(page, rows):
    pattern = re.compile(
        re.escape(ROW_BEGIN) + r".*?" + re.escape(ROW_END),
        re.S,
    )
    replacement = ROW_BEGIN + "\n" + rows.rstrip("\n") + "\n" + ROW_END
    updated, count = pattern.subn(replacement, page, count=1)
    if count != 1:
        raise SystemExit("settings page is missing BEGIN/END SETTINGS ROWS markers")
    return updated


def write_settings(rows, gen):
    SETTINGS_ROWS.parent.mkdir(parents=True, exist_ok=True)
    SETTINGS_ROWS.write_text(rows, encoding="utf-8")
    SETTINGS_GEN.write_text(gen, encoding="utf-8")
    if not SETTINGS_PAGE.is_file():
        raise SystemExit(f"missing {SETTINGS_PAGE.relative_to(ROOT)}")
    page = SETTINGS_PAGE.read_text(encoding="utf-8")
    SETTINGS_PAGE.write_text(inject_rows(page, rows), encoding="utf-8")


def _selector_decls(css, selector):
    needle = selector + " {"
    start = css.index(needle)
    end = css.index("\n}", start)
    return DECL.findall(css[start:end + 2])


def check_settings(wb, day, mobile, night, colour_day, css):
    """Return True when tokens, rows, mirrors, or settings-gen.css drift."""
    failed = False
    kept, detail = existing_tokens_kept(css)
    if not kept:
        failed = True
        print("EXISTING TOKEN BYTES CHANGED")
        print(detail)
    blocks = slice_blocks(css)
    mobile_decls = parse_decl_lines(mobile)
    if blocks["narrow"] != mobile_decls:
        failed = True
        print("560 BLOCK DOES NOT MATCH GENERATED MOBILE")
    is_mobile = _selector_decls(css, ".is-mobile")
    if is_mobile != mobile_decls:
        failed = True
        print(".is-mobile DOES NOT MATCH THE 560 BLOCK")
    dark = _selector_decls(css, ':root[data-theme="dark"]')
    scoped_dark = _selector_decls(css, '.settings-preview[data-theme="dark"]')
    if scoped_dark != dark:
        failed = True
        print("SCOPED DARK DOES NOT MATCH :root[data-theme=dark]")
    colour_decls = parse_decl_lines(colour_day)
    scoped_light = _selector_decls(css, '.settings-preview[data-theme="light"]')
    if scoped_light != colour_decls:
        failed = True
        print("SCOPED LIGHT DOES NOT MATCH THE 43 DAY COLOURS")
    root_map = dict(blocks["root"])
    xlsx_grid = {str(row.get("Token")): row for row in sheet_rows(wb, "Grid")}
    for name in NEW_GRID_ORDER:
        row = xlsx_grid.get(name)
        expected = None if row is None else css_scalar(row.get("Desktop"))
        if root_map.get(name) != expected:
            failed = True
            print(f"NEW TOKEN {name} css={root_map.get(name)!r} xlsx={expected!r}")
    for name, expected in DESKTOP_SCALE.items():
        if root_map.get(name) != expected:
            failed = True
            print(f"DESKTOP SCALE {name} css={root_map.get(name)!r} expected={expected!r}")
    if root_map.get("--card-radius") != CARD_RADIUS_HEAD:
        failed = True
        print(f"CARD RADIUS css={root_map.get('--card-radius')!r}")
    font_rows = [row for row in sheet_rows(wb, "Text") if row.get("Group") == "site-font"]
    if not font_rows:
        failed = True
        print("missing site-font rows")
    for row in font_rows:
        font_name = str(row.get("Token prefix"))
        font_expected = str(row.get("Notes") or "").strip()
        if root_map.get(font_name) != font_expected:
            failed = True
            print(f"FONT {font_name} css={root_map.get(font_name)!r} xlsx={font_expected!r}")
    for unit_name, length_name in UNITLESS_PX:
        try:
            got = float(root_map[unit_name])
            want = length_px(root_map[length_name])
        except (KeyError, ValueError) as exc:
            failed = True
            print(f"UNITLESS {unit_name} {exc}")
            continue
        if got != want:
            failed = True
            print(f"UNITLESS {unit_name}={got} length {length_name}={want}px")
    rows = settings_rows_text(wb, day, mobile)
    gen = render_settings_gen(sheet_rows(wb, "Grid"))
    current_rows = SETTINGS_ROWS.read_text(encoding="utf-8") if SETTINGS_ROWS.exists() else ""
    current_gen = SETTINGS_GEN.read_text(encoding="utf-8") if SETTINGS_GEN.exists() else ""
    if current_rows != rows:
        failed = True
        print("SETTINGS ROWS DRIFT")
        sys.stdout.writelines(difflib.unified_diff(
            current_rows.splitlines(keepends=True),
            rows.splitlines(keepends=True),
            fromfile="_shared/settings-rows.html",
            tofile="generated",
        ))
    if current_gen != gen:
        failed = True
        print("SETTINGS-GEN.CSS DRIFT")
        sys.stdout.writelines(difflib.unified_diff(
            current_gen.splitlines(keepends=True),
            gen.splitlines(keepends=True),
            fromfile="_shared/settings-gen.css",
            tofile="generated",
        ))
    if SETTINGS_PAGE.exists():
        page = SETTINGS_PAGE.read_text(encoding="utf-8")
        if inject_rows(page, rows) != page:
            failed = True
            print("SETTINGS PAGE ROWS ARE NOT THE INJECTED PARTIAL")
    else:
        failed = True
        print("MISSING site/settings/index.html")
    return failed


def command_map_lint_settings():
    """Hard fail. Formatting contexts only; displayed token values are data."""
    number_unit = re.compile(r"(?<![\w-])\d+(?:\.\d+)?(?:px|rem|em)\b")
    hex_colour = re.compile(r"#[0-9A-Fa-f]{3,8}\b")
    function_colour = re.compile(r"(?:rgb|hsl)\(")
    style_attr = re.compile(r'\sstyle="([^"]*)"')
    value = r"(?:var\(--[\w-]+\)|calc\(var\(--t-body\) \* var\(--type-scale\)\))"
    setter = re.compile(rf"^--[\w-]+:{value}(?:;--[\w-]+:{value})*$")
    failed = False
    css = SETTINGS_CSS.read_text(encoding="utf-8")
    css_hits = number_unit.findall(css) + hex_colour.findall(css) + function_colour.findall(css)
    if "@media" in css or "@container" in css or re.search(r"(?:max|min)-width:\s*\d", css):
        failed = True
        print("settings.css contains a width threshold or media query")
    if css_hits:
        failed = True
        print(f"settings.css format hits {css_hits[:20]}")
    for path in (SETTINGS_PAGE, SETTINGS_ROWS):
        text = path.read_text(encoding="utf-8")
        styles = style_attr.findall(text)
        bad = [item for item in styles if not setter.match(item)]
        stripped = style_attr.sub("", text)
        stripped = re.sub(r'<span class="token-name">(.*?)</span>', r"\1", stripped)
        stripped = re.sub(r'<span class="v\b[^"]*">.*?</span>', "", stripped, flags=re.S)
        # Workbook role prose is displayed data, same as the value column.
        stripped = re.sub(r'<div class="role">.*?</div>', "", stripped, flags=re.S)
        found = number_unit.findall(stripped) + hex_colour.findall(stripped) + function_colour.findall(stripped)
        print(
            f"{path.relative_to(ROOT)} style-attrs {len(styles)} "
            f"bad-style {len(bad)} format-outside-values {len(found)}"
        )
        if bad or found:
            failed = True
            for item in bad[:8]:
                print("BAD STYLE", item)
            for item in found[:12]:
                print("FORMAT", item)
    page = SETTINGS_PAGE.read_text(encoding="utf-8")
    if "matchMedia" in page or "innerWidth" in page or "innerHeight" in page:
        failed = True
        print("settings page has a JS width check")
    if failed:
        raise SystemExit(1)
    print("map-lint settings: 0")


def command_check():
    wb = load_workbook()
    day, mobile, night, scale, colour_day = build_parts(wb)
    css = render(day, mobile, night, colour_day)
    core = render_core(css)
    current = TOKENS.read_text(encoding="utf-8") if TOKENS.exists() else ""
    current_core = TOKENS_CORE.read_text(encoding="utf-8") if TOKENS_CORE.exists() else ""
    ok, expected, file_text = type_scale_mirror_ok(css, scale)
    failed = False
    if current != css:
        failed = True
        diff = difflib.unified_diff(
            current.splitlines(keepends=True),
            css.splitlines(keepends=True),
            fromfile="_shared/tokens.css",
            tofile="generated",
        )
        sys.stdout.writelines(diff)
    if current_core != core:
        failed = True
        diff = difflib.unified_diff(
            current_core.splitlines(keepends=True),
            core.splitlines(keepends=True),
            fromfile="_shared/tokens-core.css",
            tofile="generated",
        )
        sys.stdout.writelines(diff)
    removed = core_removals(css, core)
    if removed != list(CORE_DROP_LINES):
        failed = True
        print("tokens-core.css is not tokens.css minus exactly the 4 mobile type-scale lines")
        print(removed)
    if not ok:
        failed = True
        print("TYPE-SCALE MIRROR MISMATCH")
        if expected not in file_text:
            print("--- type-scale.css is missing the verbatim block:")
            print(expected)
        if expected not in css:
            print("--- tokens.css mirror is missing the verbatim block:")
            print(expected)
    failed = check_settings(wb, day, mobile, night, colour_day, css) or failed
    if failed:
        raise SystemExit(1)
    print("MATCH")


def command_write():
    wb = load_workbook()
    day, mobile, night, scale, colour_day = build_parts(wb)
    baseline_stop(day, mobile, night, scale)
    css = render(day, mobile, night, colour_day)
    ok, expected, file_text = type_scale_mirror_ok(css, scale)
    if not ok:
        print("STOP: refusing to change type-scale.css.")
        print("The 4-token block must stay byte-identical to main.")
        raise SystemExit(1)
    # The generator owns the block. Writing is a no-op when the bytes match,
    # which is the only legal outcome. Do not touch the file.
    if expected not in file_text:
        raise SystemExit(1)
    core = render_core(css)
    if core_removals(css, core) != list(CORE_DROP_LINES):
        print("STOP: tokens-core.css is not tokens.css minus exactly the 4 lines.")
        raise SystemExit(1)
    TOKENS.write_text(css, encoding="utf-8")
    TOKENS_CORE.write_text(core, encoding="utf-8")
    rows = settings_rows_text(wb, day, mobile)
    gen = render_settings_gen(sheet_rows(wb, "Grid"))
    write_settings(rows, gen)
    day2, mobile2, night2, _scale2, colour2 = build_parts(load_workbook())
    second = render(day2, mobile2, night2, colour2)
    if second != css or settings_rows_text(load_workbook(), day2, mobile2) != rows:
        print("STOP: generator output is not deterministic.")
        raise SystemExit(1)
    if render_core(second) != core:
        print("STOP: tokens-core.css output is not deterministic.")
        raise SystemExit(1)
    if render_settings_gen(sheet_rows(load_workbook(), "Grid")) != gen:
        print("STOP: settings-gen.css output is not deterministic.")
        raise SystemExit(1)
    print(f"written {TOKENS.relative_to(ROOT)}")
    print(f"written {TOKENS_CORE.relative_to(ROOT)}")
    print(f"written {SETTINGS_ROWS.relative_to(ROOT)}")
    print(f"written {SETTINGS_GEN.relative_to(ROOT)}")
    print(f"injected {SETTINGS_PAGE.relative_to(ROOT)}")


def command_owned():
    """Report-only. Does not gate."""
    wb = load_workbook()
    day, mobile, night, scale, colour_day = build_parts(wb)
    css = render(day, mobile, night, colour_day)
    base = slice_blocks(git_text("_shared/tokens.css"))
    base_ts = DECL.findall(git_text("_shared/type-scale.css"))
    generated_root = dict(parse_decl_lines(day))
    generated_mobile = dict(parse_decl_lines(mobile))
    generated_night = dict(parse_decl_lines(night))
    print("owned-token check (report-only)")
    print("baseline counts: "
          f":root {len(base['root'])}, "
          f"prefers-color-scheme {len(base['mq'])}, "
          f"data-theme=dark {len(base['dark'])}, "
          f"type-scale {len(base_ts)}")
    rows = []
    for name, value in base["root"]:
        rows.append((":root", name, value.strip(), generated_root.get(name)))
    for name, value in base["narrow"]:
        rows.append(("max-width-560", name, value.strip(), generated_mobile.get(name)))
    for name, value in base["mq"]:
        rows.append(("prefers-color-scheme", name, value.strip(), generated_night.get(name)))
    for name, value in base["dark"]:
        rows.append(("data-theme=dark", name, value.strip(), generated_night.get(name)))
    scale_values = dict(parse_decl_lines(scale))
    for name, value in base_ts:
        got = scale_values.get(name)
        if got is None:
            got = generated_mobile.get(name)
        rows.append(("type-scale", name, value.strip(), got))
    same = 0
    for block, name, main_value, generated in rows:
        status = "EQUAL" if generated is not None and generated.strip() == main_value else "DIFF"
        if status == "EQUAL":
            same += 1
        print(f"{status}\t{block}\t{name}\t{main_value}\t{generated}")
    pack_hits = []
    for path in sorted(ROOT.glob("*/*")):
        if not path.is_dir():
            continue
        for name in ("shared.css", "draft.css"):
            for candidate in path.rglob(name):
                if "_archive" in candidate.parts:
                    continue
                text = candidate.read_text(encoding="utf-8", errors="replace")
                count = len(re.findall(r":root\b", text))
                if count:
                    pack_hits.append(f"{candidate.relative_to(ROOT)} {count}")
    print(f"pack shared/draft :root blocks: {len(pack_hits)}")
    for hit in pack_hits:
        print(f"PACK-ROOT\t{hit}")
    print(f"baseline declarations {len(rows)}; equal {same}; diff {len(rows) - same}")
    print(f"generated mobile declarations {len(generated_mobile)}")


def command_map_lint():
    """Report-only counts. Not a failure in PR1."""
    px = re.compile(r"(?<![\w-])\d+(?:\.\d+)?px\b")
    hex_colour = re.compile(r"#[0-9A-Fa-f]{3,8}\b")
    print("map-lint (report-only, not gating in PR1)")
    totals = {"files": 0, "px": 0, "hex": 0}
    for draft in sorted(ROOT.glob("*/draft-v1-linked")):
        pack = draft.parent.name
        files = px_count = hex_count = 0
        for path in sorted(draft.rglob("*")):
            if not path.is_file() or "_archive" in path.parts:
                continue
            if path.suffix.lower() not in {".css", ".html"}:
                continue
            text = path.read_text(encoding="utf-8", errors="replace")
            files += 1
            px_count += len(px.findall(text))
            hex_count += len(hex_colour.findall(text))
        totals["files"] += files
        totals["px"] += px_count
        totals["hex"] += hex_count
        print(f"{pack}\tfiles {files}\tpx {px_count}\thex {hex_count}")
    print(f"total\tfiles {totals['files']}\tpx {totals['px']}\thex {totals['hex']}")
    print("PR1 does not fail on these counts. Hard fail is PR8.")


def owned_names(wb):
    day, mobile, night, scale, _colour_day = build_parts(wb)
    names = []
    for bucket in (day, mobile, night, scale):
        for name, _value in parse_decl_lines(bucket):
            if name not in names:
                names.append(name)
    return names


def command_repo_grep():
    wb = load_workbook()
    names = owned_names(wb)
    print("repo grep: other files that define a tokens.css-owned name (report-only)")
    print(f"owned names {len(names)}")
    hits = []
    for path in sorted(ROOT.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in {".css", ".html"}:
            continue
        rel = path.relative_to(ROOT).as_posix()
        if rel == "_shared/tokens.css":
            continue
        if "_archive" in path.parts:
            continue
        text = path.read_text(encoding="utf-8", errors="replace")
        found = sorted(set(PROP.findall(text)))
        owned = [name for name in found if name in names]
        if not owned:
            continue
        hits.append(rel)
        print(f"{rel}\t{len(owned)}\t{', '.join(owned)}")
    print(f"files {len(hits)}")


def _strip_css_comments(text):
    return re.sub(r"/\*.*?\*/", "", text, flags=re.S)


def _strip_at_preludes(text):
    """Drop @media / @container / @supports conditions. Bodies stay."""
    out = []
    i = 0
    names = ("@media", "@container", "@supports")
    while i < len(text):
        if text.startswith(names, i):
            brace = text.find("{", i)
            if brace < 0:
                out.append(text[i])
                i += 1
                continue
            i = brace
            continue
        out.append(text[i])
        i += 1
    return "".join(out)


_LENGTH_HIT = re.compile(r"(?<![\w.-])(-?(?:\d+\.\d+|\d+|\.\d+))(px|rem|em)\b")
_HEX_HIT = re.compile(r"#[0-9A-Fa-f]{3,8}\b")
_FN_HIT = re.compile(r"(?:rgba?|hsla?)\(")
_FW_HIT = re.compile(r"font-weight\s*:\s*(\d+)\b")
_LH_HIT = re.compile(
    r"(?:(?<=^)|(?<=[;{]))\s*line-height\s*:\s*(-?(?:\d+\.\d+|\d+|\.\d+))\s*(?:!important)?\s*(?=;)",
    re.M,
)
# Q-MAP-4 freezes cite-1.css apart from --cite-font, so em/weight/line-height
# that already live in untouched see-also.css are not part of the shared gate.
# px, rem, hex and colour functions are gated on every _shared file except the
# two generated token files. The stricter set is gated on every other file.
_LINT_LOOSE = {
    "_shared/see-also.css",
    "_shared/settings.css",
    "_shared/settings-gen.css",
}
_F1_EXEMPT = tuple(item.strip() for item in CORE_DROP_LINES)


def _blank_calls(text):
    """Hide var() and url() interiors. Fallbacks such as 61.25rem stay put."""
    chars = list(text)
    i = 0
    n = len(text)
    while i < n:
        if text.startswith(("var(", "url("), i):
            depth = 0
            j = text.find("(", i)
            while j < n:
                if text[j] == "(":
                    depth += 1
                elif text[j] == ")":
                    depth -= 1
                    if depth == 0:
                        j += 1
                        break
                j += 1
            for k in range(i, min(j, n)):
                if chars[k] != "\n":
                    chars[k] = " "
            i = j
            continue
        i += 1
    return "".join(chars)


def shared_lint_hits(text, strict):
    cleaned = _strip_at_preludes(_strip_css_comments(text))
    kept = []
    for line in cleaned.splitlines():
        if line.strip() in _F1_EXEMPT:
            continue
        kept.append(line)
    body = _blank_calls("\n".join(kept))
    hits = []
    for match in _LENGTH_HIT.finditer(body):
        if float(match.group(1)) == 0:
            continue
        unit = match.group(2)
        if unit == "em" and not strict:
            continue
        hits.append(match.group(0))
    hits.extend(_HEX_HIT.findall(body))
    hits.extend(_FN_HIT.findall(body))
    if strict:
        hits.extend(f"font-weight:{item}" for item in _FW_HIT.findall(body))
        for match in _LH_HIT.finditer(body):
            if float(match.group(1)) == 0:
                continue
            hits.append("line-height:" + match.group(1))
    return hits


def command_map_lint_shared():
    """Hard fail. tokens.css and tokens-core.css are generated and excluded."""
    failed = False
    root = ROOT / "_shared"
    files = sorted(path for path in root.rglob("*.css") if path.is_file())
    for path in files:
        rel = path.relative_to(ROOT).as_posix()
        if rel in ("_shared/tokens.css", "_shared/tokens-core.css"):
            print(f"{rel}\texcluded")
            continue
        hits = shared_lint_hits(path.read_text(encoding="utf-8"), strict=rel not in _LINT_LOOSE)
        print(f"{rel}\t{len(hits)}")
        if hits:
            failed = True
            for item in hits[:12]:
                print(f"  {item}")
    if failed:
        raise SystemExit(1)
    print("map-lint shared: 0")


def command_settings_cover():
    """Pre-existing token names stay covered. PR3 additions are exempt.

    check-settings-rows.py is unchanged. This mode is the CI gate so new
    LEGACY-PR3 and font/border rows do not force a settings HTML edit.
    """
    page = SETTINGS_PAGE.read_text(encoding="utf-8")
    partial = SETTINGS_ROWS.read_text(encoding="utf-8")
    start = page.index(ROW_BEGIN) + len(ROW_BEGIN)
    end = page.index(ROW_END)
    injected = page[start:end].strip("\n")
    if injected != partial.strip("\n"):
        print("settings page rows differ from _shared/settings-rows.html")
        raise SystemExit(1)
    row_re = re.compile(r'<tr class="row\b[^"]*">.*?</tr>', re.S)
    code_re = re.compile(r"<code>(--[\w-]+)</code>")
    rows = row_re.findall(page)
    counts = {}
    for row in rows:
        for name in code_re.findall(row):
            counts[name] = counts.get(name, 0) + 1
    defined = []
    seen = set()
    for path in (TOKENS, TYPE_SCALE):
        for name in PROP.findall(path.read_text(encoding="utf-8")):
            if name not in seen:
                seen.add(name)
                defined.append(name)

    def exempt(name):
        return name.startswith("--legacy-") or name in SETTINGS_HIDE_NAMES

    missing = [name for name in defined if counts.get(name, 0) == 0 and not exempt(name)]
    duplicated = [name for name, count in counts.items() if count != 1]
    extra = [name for name in counts if name not in seen]
    placeholders = all(title in page for title in ("Box", "Components", "TopBar"))
    covered = len([name for name in defined if not exempt(name)]) - len(missing)
    print(
        f"{len(defined)} tokens / {covered} covered / {len(missing)} missing / "
        f"{len(duplicated)} duplicated / {len(rows)} rows"
    )
    if missing:
        print("missing: " + ", ".join(missing))
    if duplicated:
        print("duplicated: " + ", ".join(duplicated))
    if extra:
        print("not a token: " + ", ".join(extra))
    if not placeholders:
        print("missing Box, Components, or TopBar placeholder")
    if missing or duplicated or extra or not placeholders:
        raise SystemExit(1)
    print("settings-cover: ok")


def main():
    args = set(sys.argv[1:])
    if args == {"--check"}:
        command_check()
        return
    if args == {"--owned"}:
        command_owned()
        return
    if args == {"--map-lint"}:
        command_map_lint()
        return
    if args == {"--map-lint-shared"}:
        command_map_lint_shared()
        return
    if args == {"--settings-cover"}:
        command_settings_cover()
        return
    if args == {"--repo-grep"}:
        command_repo_grep()
        return
    if args == {"--settings"}:
        command_write()
        return
    if args == {"--map-lint-settings"}:
        command_map_lint_settings()
        return
    if args:
        raise SystemExit(
            "usage: xlsx_to_tokens.py [--check|--owned|--map-lint|--map-lint-shared|"
            "--map-lint-settings|--settings-cover|--repo-grep|--settings]"
        )
    command_write()


if __name__ == "__main__":
    main()
