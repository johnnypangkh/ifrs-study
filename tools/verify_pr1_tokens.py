#!/usr/bin/env python3
"""Verify every token in main _shared/tokens.css + type-scale.css has an xlsx row reproducing its value exactly.

Paths are the repo copies: design/MASTER-STYLES.xlsx, and main's
_shared/tokens.css + _shared/type-scale.css (origin/main, else main).
The working tree tokens.css is the generated file and carries the new
tokens as well; this check stays on the 146 main declarations.
"""
import re
import subprocess
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "design" / "MASTER-STYLES.xlsx"


def git_text(rel):
    import os
    refs = ["origin/main", "main"]
    base = os.environ.get("GITHUB_BASE_REF", "")
    if base:
        refs.extend([f"origin/{base}", base])
    for ref in refs:
        try:
            return subprocess.check_output(
                ["git", "show", f"{ref}:{rel}"],
                cwd=ROOT,
                text=True,
                stderr=subprocess.DEVNULL,
            )
        except subprocess.CalledProcessError:
            continue
    raise SystemExit(f"cannot read {rel} from main")


css = git_text("_shared/tokens.css")
ts = git_text("_shared/type-scale.css")
a = css.index(":root {")
b = css.index("@media (prefers-color-scheme")
c = css.index(':root[data-theme="dark"]')
blk = lambda s: re.findall(r"(--[\w-]+)\s*:\s*([^;]+);", s)
main = {"root": blk(css[a:b]), "mq": blk(css[b:c]), "dark": blk(css[c:]), "ts": blk(ts)}
wb = openpyxl.load_workbook(XLSX, data_only=True)
n = lambda v: (("%g" % float(v)).lstrip("0") if float(v) < 1 else "%g" % float(v))
gen = {"root": {}, "dark": {}, "ts": {}}
for r in wb["Colour"].iter_rows(min_row=2, values_only=True):
    gen["root"][r[1]] = r[2]
    if r[3] != "(same as day)":
        gen["dark"][r[1]] = r[3]
for r in wb["Text"].iter_rows(min_row=2, values_only=True):
    if r[-1] == "site":
        gen["root"][r[1]] = n(r[3]) + "rem"
    if r[-1] == "type-scale":
        notes = "" if r[10] is None else str(r[10])
        if "mobile " in notes:
            gen["ts"][r[1]] = notes.split("mobile ", 1)[1].strip()
        elif len(r) > 6 and r[6] is not None and str(r[6]).strip():
            gen["ts"][r[1]] = str(r[6]).strip()
for r in wb["Spacing"].iter_rows(min_row=2, values_only=True):
    gen["root"][r[0]] = r[1]
for r in wb["Grid"].iter_rows(min_row=2, values_only=True):
    if r[1] == "--page-max":
        gen["root"][r[1]] = r[2]
bad = 0
tot = 0
for blkn, src in (("root", "root"), ("mq", "dark"), ("dark", "dark"), ("ts", "ts")):
    for k, v in main[blkn]:
        tot += 1
        g = gen[src].get(k)
        if g is None or str(g).strip() != v.strip():
            bad += 1
            print("MISMATCH", blkn, k, "main=", v, "xlsx=", g)
print(f"{tot} declarations checked, {bad} mismatches")
sys.exit(bad > 0)
